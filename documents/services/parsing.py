import csv
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd


def detect_csv_options(path: Path) -> tuple[str, str]:
    """Detecta la codificación y el delimitador de un archivo CSV."""
    raw = path.read_bytes()[:65536]
    if not raw:
        return "utf-8", ","
    if raw.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    else:
        try:
            raw.decode("utf-8")
            encoding = "utf-8"
        except UnicodeDecodeError:
            encoding = "latin-1"

    text = raw.decode(encoding, errors="replace")[:4096]
    delims = [",", ";", "\t", "|"]
    try:
        sniffer = csv.Sniffer()
        dialect = sniffer.sniff(text, delimiters=delims)
        delimiter = dialect.delimiter
    except Exception:
        counts = {c: text.count(c) for c in delims}
        delimiter = max(counts, key=counts.get) if counts else ","
        if counts.get(delimiter, 0) == 0:
            delimiter = ","
    return encoding, delimiter


def normalize_headers(raw_headers) -> list[str]:
    """Normaliza nombres de columnas vacíos o duplicados sin perder datos."""
    clean = []
    seen = {}
    for i, h in enumerate(raw_headers):
        val = str(h).strip() if h is not None else ""
        if "." in val and val.rsplit(".", 1)[-1].isdigit():
            val = val.rsplit(".", 1)[0]
        if not val or val.lower().startswith("unnamed:"):
            val = f"Columna_{i + 1}"
        base = val
        count = seen.get(base.lower(), 0)
        seen[base.lower()] = count + 1
        if count > 0:
            val = f"{base}_{count + 1}"
        clean.append(val)
    return clean


def parse_smart_val(v):
    """
    Conserva identificadores con ceros a la izquierda como string ('00123'),
    convierte enteros y decimales a tipos numéricos, y mantiene textos/fechas.
    """
    if v is None or pd.isna(v):
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if not s or s.lower() == "nan":
        return None
    if len(s) > 1 and s.startswith("0") and s.isdigit():
        return s
    try:
        if "." not in s and "e" not in s.lower():
            return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        pass
    return s


def detect_excel_header_row(rows: list) -> int:
    """Busca la primera fila que probablemente contiene los encabezados reales."""
    if not rows:
        return 0
    counts = [sum(1 for c in r if c is not None and str(c).strip() != "") for r in rows]
    max_c = max(counts) if counts else 0
    if max_c <= 1:
        return 0
    for idx, c in enumerate(counts):
        if c >= max(2, int(max_c * 0.6)):
            return idx
    return 0


def read_sheets(path: Path, options: Optional[dict] = None) -> tuple[dict[str, pd.DataFrame], dict]:
    """Devuelve ({nombre_hoja: DataFrame}, import_options_detectadas)."""
    options = options or {}
    size = path.stat().st_size
    if size == 0:
        raise ValueError("El archivo está vacío (0 bytes).")

    ext = path.suffix.lower()
    used_options = {}

    if ext == ".csv":
        enc = options.get("encoding")
        sep = options.get("delimiter")
        auto_enc, auto_sep = detect_csv_options(path)
        encoding = enc or auto_enc
        delimiter = sep or auto_sep
        used_options = {"encoding": encoding, "delimiter": delimiter}
        try:
            df = pd.read_csv(path, sep=delimiter, encoding=encoding, dtype=object, keep_default_na=False)
        except UnicodeDecodeError:
            df = pd.read_csv(path, sep=delimiter, encoding="latin-1", dtype=object, keep_default_na=False)
            used_options["encoding"] = "latin-1"
        except Exception as e:
            raise ValueError(f"No se pudo leer el archivo CSV: {e}")

        if not df.empty and df.columns is not None:
            df.columns = normalize_headers(df.columns)
            df = df.map(parse_smart_val)
        return {"Datos": df}, used_options

    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        sheet_names = wb.sheetnames
    except Exception:
        try:
            excel_file = pd.ExcelFile(path)
            sheet_names = excel_file.sheet_names
            wb = None
        except Exception:
            raise ValueError("El archivo no es un libro de Excel válido o está dañado.")

    sheets = {}
    excel_headers = {}
    for sname in sheet_names:
        h_idx = options.get(f"header_row_{sname}", options.get("header_row"))
        if h_idx is None:
            if wb is not None:
                try:
                    ws = wb[sname]
                    rows = list(ws.iter_rows(values_only=True, max_row=15))
                    counts = [sum(1 for c in r if c is not None and str(c).strip() != "") for r in rows]
                    max_c = max(counts) if counts else 0
                    if max_c == 0:
                        sheets[sname] = pd.DataFrame()
                        continue
                    h_idx = detect_excel_header_row(rows)
                except Exception:
                    h_idx = 0
            else:
                h_idx = 0
        else:
            h_idx = int(h_idx)

        excel_headers[sname] = h_idx
        try:
            df = pd.read_excel(path, sheet_name=sname, header=h_idx, dtype=object)
            if not df.empty and df.columns is not None:
                df.columns = normalize_headers(df.columns)
                df = df.map(parse_smart_val)
            sheets[sname] = df
        except Exception:
            sheets[sname] = pd.DataFrame()

    if wb is not None:
        wb.close()
    used_options = {"headers": excel_headers}
    return sheets, used_options


def infer_schema_type_info(values) -> dict:
    """Infers display metadata without coercing or changing stored cell values."""
    kinds = set()
    for value in values:
        if value is None or pd.isna(value):
            continue
        if isinstance(value, bool):
            kinds.add("boolean")
        elif isinstance(value, (int, float)):
            kinds.add("number")
        elif isinstance(value, (datetime, pd.Timestamp)):
            kinds.add("date")
        elif isinstance(value, str):
            text = value.strip()
            if text:
                if len(text) >= 10 and text[4:5] == "-" and text[7:8] == "-" and text[:4].isdigit():
                    kinds.add("date")
                else:
                    kinds.add("text")
        else:
            kinds.add(type(value).__name__)

    if not kinds:
        return {"type": "empty", "observed_types": []}
    ordered_kinds = [kind for kind in ("boolean", "date", "number", "text") if kind in kinds]
    ordered_kinds.extend(sorted(kinds - set(ordered_kinds)))
    return {
        "type": ordered_kinds[0] if len(ordered_kinds) == 1 else "mixed",
        "observed_types": ordered_kinds,
    }


def infer_schema_value_type(values) -> str:
    return infer_schema_type_info(values)["type"]


def build_schema_profile(sheets: dict[str, pd.DataFrame]) -> dict:
    return {
        "profile_version": 1,
        "sheets": [
            {
                "name": str(sheet_name),
                "columns": [
                    {
                        "name": str(column),
                        **infer_schema_type_info(df[column]),
                    }
                    for column in df.columns
                ],
            }
            for sheet_name, df in sheets.items()
        ],
    }


def compare_schema_profiles(previous: dict, current: dict) -> list[dict]:
    """Returns explicit structural changes; single add/remove pairs are only rename candidates."""
    changes = []
    previous_sheets = {sheet["name"]: sheet for sheet in previous.get("sheets", [])}
    current_sheets = {sheet["name"]: sheet for sheet in current.get("sheets", [])}

    for name in sorted(current_sheets.keys() - previous_sheets.keys()):
        changes.append({"kind": "sheet_added", "sheet": name})
    for name in sorted(previous_sheets.keys() - current_sheets.keys()):
        changes.append({"kind": "sheet_removed", "sheet": name})

    for sheet_name in sorted(previous_sheets.keys() & current_sheets.keys()):
        old_columns = previous_sheets[sheet_name].get("columns", [])
        new_columns = current_sheets[sheet_name].get("columns", [])
        old_by_name = {column["name"]: (index, column) for index, column in enumerate(old_columns)}
        new_by_name = {column["name"]: (index, column) for index, column in enumerate(new_columns)}
        removed = [column for column in old_columns if column["name"] not in new_by_name]
        added = [column for column in new_columns if column["name"] not in old_by_name]

        rename_candidates = []
        if len(removed) == 1 and len(added) == 1 and removed[0]["type"] == added[0]["type"]:
            rename_candidates = [(removed[0], added[0])]
            changes.append({
                "kind": "possible_column_rename",
                "sheet": sheet_name,
                "from": removed[0]["name"],
                "to": added[0]["name"],
            })

        for column in removed:
            if not rename_candidates:
                changes.append({"kind": "column_removed", "sheet": sheet_name, "column": column["name"]})
        for column in added:
            if not rename_candidates:
                changes.append({"kind": "column_added", "sheet": sheet_name, "column": column["name"]})

        for name in sorted(old_by_name.keys() & new_by_name.keys()):
            old_index, old_column = old_by_name[name]
            new_index, new_column = new_by_name[name]
            if old_column["type"] != new_column["type"]:
                changes.append({
                    "kind": "column_type_changed",
                    "sheet": sheet_name,
                    "column": name,
                    "from": old_column["type"],
                    "to": new_column["type"],
                })
        old_order = [column["name"] for column in old_columns if column["name"] in new_by_name]
        new_order = [column["name"] for column in new_columns if column["name"] in old_by_name]
        if old_order != new_order:
            changes.append({"kind": "column_order_changed", "sheet": sheet_name})

    return changes


