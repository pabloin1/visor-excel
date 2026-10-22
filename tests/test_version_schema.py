import io
import json
import tempfile
from datetime import datetime
from pathlib import Path

from django.test import TestCase, override_settings
from openpyxl import Workbook

from documents.models import Version
from documents.services import history as app
from documents.services.history import HTTPException


def make_workbook(headers, rows, sheet_name="Datos"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def make_empty_workbook(sheet_name="Vacia"):
    workbook = Workbook()
    workbook.active.title = sheet_name
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class VersionSchemaTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.settings_override = override_settings(
            MEDIA_ROOT=self.root / "data" / "uploads",
            VISOR_CACHE_ROOT=self.root / "data" / "cache",
        )
        self.settings_override.enable()

    def tearDown(self):
        self.settings_override.disable()
        self.temp_dir.cleanup()

    def upload(self, filename, content, mode="", comparison_keys=None):
        return app.upload(
            filename=filename,
            content=content,
            mode=mode,
            comparison_keys=(
                json.dumps(comparison_keys) if comparison_keys is not None else None
            ),
        )

    def test_same_structure_different_values_has_no_schema_changes(self):
        first = make_workbook(["Clave", "Monto"], [["0012", 10]])
        second = make_workbook(["Clave", "Monto"], [["0034", 500]])
        self.upload("reporte.xlsx", first)

        result = self.upload("reporte.xlsx", second)

        self.assertEqual(result["status"], "name_match")
        self.assertFalse(result["schema_comparison"]["has_changes"])
        self.assertEqual(result["schema_comparison"]["changes"], [])

    def test_single_same_type_column_replacement_is_only_a_rename_candidate(self):
        original = make_workbook(["Beneficiario", "Monto"], [["Ana", 10]])
        revised = make_workbook(["Persona", "Monto"], [["Ana", 20]])
        self.upload("reporte.xlsx", original)

        comparison = self.upload("reporte.xlsx", revised)["schema_comparison"]

        self.assertTrue(comparison["has_changes"])
        self.assertIn(
            {
                "kind": "possible_column_rename",
                "sheet": "Datos",
                "from": "Beneficiario",
                "to": "Persona",
            },
            comparison["changes"],
        )

    def test_structural_changes_are_reported_and_saved_on_new_version(self):
        first = make_workbook(["Clave", "Monto"], [["0012", 10]])
        second = make_workbook(["Clave", "Monto", "Estado"], [["0034", "diez", "Activo"]])
        self.upload("reporte.xlsx", first)

        comparison = self.upload("reporte.xlsx", second)

        self.assertTrue(comparison["schema_comparison"]["has_changes"])
        changes = comparison["schema_comparison"]["changes"]
        self.assertIn(
            {"kind": "column_added", "sheet": "Datos", "column": "Estado"},
            changes,
        )
        self.assertIn(
            {
                "kind": "column_type_changed",
                "sheet": "Datos",
                "column": "Monto",
                "from": "number",
                "to": "text",
            },
            changes,
        )

        setup = self.upload("reporte.xlsx", second, mode="new_version")
        self.assertEqual(setup["status"], "comparison_setup_required")
        saved = self.upload(
            "reporte.xlsx",
            second,
            mode="new_version",
            comparison_keys={"Datos": ["Clave"]},
        )
        profile = Version.objects.get(pk=saved["version_id"]).schema_info
        self.assertEqual(
            [column["name"] for column in profile["sheets"][0]["columns"]],
            ["Clave", "Monto", "Estado"],
        )

    def test_legacy_profile_is_reconstructed_before_comparison(self):
        first = make_workbook(["Clave", "Monto"], [["0012", 10]])
        second = make_workbook(["Clave", "Monto"], [["0034", 500]])
        original = self.upload("reporte.xlsx", first)
        Version.objects.filter(pk=original["version_id"]).update(schema_info=None)

        result = self.upload("reporte.xlsx", second)

        self.assertFalse(result["schema_comparison"]["has_changes"])
        stored = Version.objects.get(pk=original["version_id"]).schema_info
        self.assertIsNotNone(stored)

    def test_invalid_candidate_does_not_create_version(self):
        first = make_workbook(["Clave"], [["0012"]])
        self.upload("reporte.xlsx", first)

        with self.assertRaises(HTTPException) as error:
            self.upload("reporte.xlsx", b"not an xlsx file")

        self.assertEqual(error.exception.status_code, 400)
        count = Version.objects.count()
        self.assertEqual(count, 1)

    def test_mixed_column_types_are_exposed_and_numeric_summaries_report_omissions(self):
        workbook = make_workbook(
            ["Monto", "Referencia"],
            [[10, "A-001"], ["Pendiente", "A-002"], [None, "A-003"]],
        )
        uploaded = self.upload("mixto.xlsx", workbook)

        version = app.get_version(uploaded["version_id"])
        sheet = version["sheets"][0]
        self.assertEqual(sheet["column_types"][0]["type"], "mixed")
        self.assertEqual(sheet["column_types"][0]["observed_types"], ["number", "text"])

        summary = app.get_version_summary(
            uploaded["version_id"],
            sheet="Datos",
            measure_col=0,
            stat_op="sum",
        )
        self.assertIn(["Columnas mixtas", 1], summary["metrics"])
        self.assertEqual(summary["chart"]["omitted_non_numeric"], 1)
        self.assertEqual(summary["chart"]["omitted_nulls"], 1)

        grouped = app.get_version_summary(
            uploaded["version_id"],
            sheet="Datos",
            category_col=0,
            stat_op="count",
        )
        self.assertEqual(
            {label: value for label, value in grouped["chart"]["values"]},
            {"10.0": 1, "Pendiente": 1, "(Sin dato)": 1},
        )

    def test_empty_sheet_returns_empty_rows_and_summary_without_filter_messaging(self):
        uploaded = self.upload("vacio.xlsx", make_empty_workbook())

        version = app.get_version(uploaded["version_id"])
        self.assertEqual(version["sheets"][0]["row_count"], 0)
        self.assertEqual(version["sheets"][0]["columns"], [])

        rows = app.get_version_rows(uploaded["version_id"], sheet="Vacia")
        self.assertEqual(rows["total_rows"], 0)
        self.assertEqual(rows["rows"], [])

        summary = app.get_version_summary(uploaded["version_id"], sheet="Vacia")
        self.assertEqual(summary["chart"]["values"], [])
        self.assertEqual(summary["chart"]["title"], "Sin datos en la hoja")

    def test_excel_workbook_report_profiles_sheets_and_skips_identifiers(self):
        book = Workbook()
        main = book.active
        main.title = "Ventas"
        main.append(["Folio", "Programa", "Monto", "Fecha"])
        main.append([1001, "A", 10, datetime(2026, 1, 2)])
        main.append([1002, "A", "Pendiente", datetime(2026, 1, 3)])
        main.append([1003, "B", None, None])
        empty = book.create_sheet("Sin movimientos")
        empty.append(["Referencia", "Estado"])
        stream = io.BytesIO()
        book.save(stream)
        uploaded = self.upload("reporte_excel.xlsx", stream.getvalue())
        report = app.get_workbook_report(uploaded["version_id"])

        self.assertEqual(report["sheet_count"], 2)
        self.assertEqual(report["total_rows"], 3)
        self.assertEqual(report["empty_sheets"], 1)
        self.assertEqual(report["sheets_with_observations"], 2)
        ventas = report["sheets"][0]
        self.assertEqual(ventas["rows_with_missing"], 1)
        self.assertEqual(ventas["mixed_columns"], 1)
        self.assertEqual(ventas["suggested_category"], 1)
        self.assertEqual(ventas["suggested_measure"], 2)
        self.assertTrue(ventas["numeric_fields"][0]["identifier"])
        self.assertNotIn("sum", ventas["numeric_fields"][0])
        amount = ventas["numeric_fields"][1]
        self.assertEqual((amount["valid"], amount["non_numeric"], amount["missing"]), (1, 1, 1))
        self.assertEqual(amount["sum"], 10)
        self.assertEqual((ventas["date_fields"][0]["minimum"], ventas["date_fields"][0]["maximum"]),
                         ("2026-01-02", "2026-01-03"))
        self.assertEqual(report["sheets"][1]["columns"], 2)

        count = app.get_version_summary(uploaded["version_id"], "Ventas", category_col=1,
                                        measure_col=2, stat_op="count")
        self.assertEqual(dict(count["chart"]["values"]), {"A": 2, "B": 1})
        self.assertEqual((count["chart"]["omitted_nulls"], count["chart"]["omitted_non_numeric"]), (0, 0))
        all_rows = app.get_version_summary(uploaded["version_id"], "Ventas", measure_col=2, stat_op="count")
        self.assertEqual(all_rows["chart"]["values"], [["Todos los registros", 3]])

    def test_other_average_is_weighted_by_valid_numeric_records(self):
        rows = [[letter, value] for letter, value in zip("ABCDEFG", range(100, 93, -1))]
        rows += [["H", 1], ["H", 1], ["H", 1], ["I", 3]]
        uploaded = self.upload("promedios.xlsx", make_workbook(["Grupo", "Monto"], rows))
        result = app.get_version_summary(uploaded["version_id"], "Datos", category_col=0,
                                         measure_col=1, stat_op="avg")
        self.assertEqual(result["chart"]["values"][-1], ["Otros", 1.5])
        self.assertEqual(len(result["chart"]["other_details"]), 2)
        compact = app.get_version_summary(uploaded["version_id"], "Datos", category_col=0,
                                          measure_col=1, stat_op="avg", chart_page=0, other_page=0)
        self.assertEqual(compact["chart"]["values"][-1], ["Otros", 1.5])
        self.assertEqual(compact["chart"]["other_total"], 2)

        no_measure = app.get_version_summary(uploaded["version_id"], "Datos", category_col=0,
                                             stat_op="avg", chart_page=0, other_page=0)
        self.assertEqual(no_measure["chart"]["title"], "Recuento de registros por Grupo")
        self.assertEqual(no_measure["chart"]["values"][-1], ["Otros", 2.0])

    def test_unique_amounts_remain_measures(self):
        uploaded = self.upload("montos.xlsx", make_workbook(
            ["Folio", "Monto"], [[f"F-{i:03d}", i * 7] for i in range(1, 13)]
        ))
        sheet = app.get_workbook_report(uploaded["version_id"])["sheets"][0]
        self.assertEqual(sheet["suggested_measure"], 1)
        self.assertFalse(sheet["numeric_fields"][0]["identifier"])
        self.assertEqual(sheet["numeric_fields"][0]["valid"], 12)

    def test_upload_duplicate_rows_summary_and_delete_regression(self):
        workbook = make_workbook(
            ["Grupo", "Monto"],
            [["A", 1], ["B", 3], ["A", 2], ["B", 4]],
        )
        uploaded = self.upload("ventas.xlsx", workbook)
        duplicate = self.upload("ventas.xlsx", workbook)
        self.assertEqual(duplicate, {"status": "duplicate", "version_id": uploaded["version_id"]})

        version = app.get_version(uploaded["version_id"])
        self.assertEqual(version["sheets"][0]["name"], "Datos")
        page = app.get_version_rows(
            uploaded["version_id"],
            sheet="Datos",
            page=0,
            page_size=1,
            q="a",
            sort_col=1,
            sort_dir="desc",
            filters=json.dumps({"0": "A"}),
        )
        self.assertEqual(page["filtered_rows"], 2)
        self.assertEqual(page["rows"][0]["r"], ["A", 2])

        no_matches = app.get_version_rows(
            uploaded["version_id"], sheet="Datos", q="sin-coincidencias"
        )
        self.assertGreater(no_matches["total_rows"], 0)
        self.assertEqual(no_matches["filtered_rows"], 0)
        self.assertEqual(no_matches["rows"], [])

        summary = app.get_version_summary(
            uploaded["version_id"],
            sheet="Datos",
            category_col=0,
            measure_col=1,
            stat_op="sum",
        )
        self.assertEqual(
            {label: value for label, value in summary["chart"]["values"]},
            {"A": 3, "B": 7},
        )

        file_meta = Version.objects.get(pk=uploaded["version_id"])
        stored_path = app.uploads_dir() / file_meta.stored_path
        cache_path = app.cache_dir() / f"{file_meta.sha256}.db"
        app.delete_version(uploaded["version_id"])
        self.assertFalse(stored_path.exists())
        self.assertFalse(cache_path.exists())
        with self.assertRaises(HTTPException) as error:
            app.get_version(uploaded["version_id"])
        self.assertEqual(error.exception.status_code, 404)

    def test_representative_fixtures_preserve_imported_data_and_options(self):
        fixture_dir = Path(__file__).resolve().parent.parent / "muestras" / "casos_prueba"

        comma_csv = self.upload(
            "comma.csv", (fixture_dir / "04_csv_coma_utf8.csv").read_bytes()
        )
        comma_version = app.get_version(comma_csv["version_id"])
        self.assertEqual(comma_version["sheets"][0]["row_count"], 4)
        self.assertIn("Café_Quintales", comma_version["sheets"][0]["columns"])

        latin1 = self.upload(
            "latin1.csv", (fixture_dir / "05_csv_puntoycoma_latin1.csv").read_bytes()
        )
        latin1_version = app.get_version(latin1["version_id"])
        self.assertEqual(latin1_version["import_options"]["encoding"], "latin-1")
        self.assertEqual(latin1_version["import_options"]["delimiter"], ";")
        self.assertEqual(
            latin1_version["sheets"][0]["columns"],
            ["Identificador;Folio", "Descripción_Operación", "Monto_Ejecutado"],
        )

        tab_csv = self.upload(
            "tab.csv", (fixture_dir / "06_csv_tab_utf8_sig.csv").read_bytes()
        )
        tab_version = app.get_version(tab_csv["version_id"])
        self.assertEqual(tab_version["import_options"]["encoding"], "utf-8-sig")
        self.assertEqual(tab_version["import_options"]["delimiter"], "\t")
        self.assertEqual(len(tab_version["sheets"][0]["columns"]), 3)

        mixed = self.upload(
            "mixed.xlsx", (fixture_dir / "03_excel_tipos_mixtos_y_ceros.xlsx").read_bytes()
        )
        mixed_version = app.get_version(mixed["version_id"])
        self.assertEqual(mixed_version["sheets"][0]["column_types"][3]["type"], "mixed")
        first_row = app.get_version_rows(
            mixed["version_id"], sheet="Padron", page=0, page_size=1
        )["rows"][0]["r"]
        self.assertEqual(first_row[:2], ["00129482", "02840"])

        decorated = self.upload(
            "decorated.xlsx",
            (fixture_dir / "01_excel_metadatos_previos.xlsx").read_bytes(),
        )
        decorated_version = app.get_version(decorated["version_id"])
        self.assertEqual(
            decorated_version["sheets"][0]["columns"],
            ["Codigo", "Nombre", "Departamento", "Presupuesto"],
        )
        self.assertEqual(
            decorated_version["import_options"]["headers"]["Reporte"], 3
        )

    def test_edge_case_fixtures_are_handled_without_partial_versions(self):
        fixture_dir = Path(__file__).resolve().parent.parent / "muestras" / "casos_prueba"

        duplicate_headers = self.upload(
            "headers.xlsx",
            (fixture_dir / "02_excel_encabezados_duplicados_y_vacios.xlsx").read_bytes(),
        )
        headers = app.get_version(duplicate_headers["version_id"])["sheets"][0]["columns"]
        self.assertEqual(
            headers,
            ["ID", "Descripcion", "Columna_3", "Cantidad", "Cantidad_2"],
        )

        empty_sheets = self.upload(
            "empty-sheets.xlsx",
            (fixture_dir / "08_excel_hojas_vacias.xlsx").read_bytes(),
        )
        sheets = app.get_version(empty_sheets["version_id"])["sheets"]
        self.assertEqual([sheet["row_count"] for sheet in sheets], [2, 0, 0])

        for filename in (
            "07_archivo_vacio.xlsx",
            "07_archivo_vacio.csv",
            "09_archivo_corrupto.xlsx",
        ):
            with self.subTest(filename=filename):
                with self.assertRaises(HTTPException) as error:
                    self.upload(filename, (fixture_dir / filename).read_bytes())
                self.assertEqual(error.exception.status_code, 400)

        version_count = Version.objects.count()
        self.assertEqual(version_count, 2)

    def test_large_csv_is_stored_and_queried_with_server_pagination(self):
        fixture_dir = Path(__file__).resolve().parent.parent / "muestras" / "casos_prueba"
        uploaded = self.upload(
            "large.csv", (fixture_dir / "10_archivo_grande.csv").read_bytes()
        )

        version = app.get_version(uploaded["version_id"])
        self.assertEqual(version["sheets"][0]["row_count"], 35000)

        page = app.get_version_rows(
            uploaded["version_id"],
            sheet="Datos",
            page=0,
            page_size=25,
            q="selva",
            filters=json.dumps({"3": "Fertilizante"}),
            sort_col=6,
            sort_dir="desc",
        )
        self.assertEqual(page["total_rows"], 35000)
        self.assertLessEqual(len(page["rows"]), 25)
        self.assertGreater(page["filtered_rows"], 0)
        self.assertTrue(all(row["r"][2] == "Selva" for row in page["rows"]))
        self.assertTrue(all(row["r"][3] == "Fertilizante" for row in page["rows"]))

    def test_summary_collapses_many_categories_and_wide_sheets_keep_all_columns(self):
        grouped = self.upload(
            "categorias.xlsx",
            make_workbook(
                ["Categoria", "Valor"],
                [[f"Categoria {index}", index] for index in range(12)],
            ),
        )
        summary = app.get_version_summary(
            grouped["version_id"],
            sheet="Datos",
            category_col=0,
            measure_col=1,
            stat_op="sum",
        )
        self.assertEqual(len(summary["chart"]["values"]), 8)
        self.assertEqual(len(summary["chart"]["other_details"]), 5)

        columns = [f"Campo {index}" for index in range(15)]
        wide = self.upload("ancho.xlsx", make_workbook(columns, [[str(i) for i in range(15)]]))
        wide_version = app.get_version(wide["version_id"])
        self.assertEqual(len(wide_version["sheets"][0]["columns"]), 15)
        wide_row = app.get_version_rows(wide["version_id"], sheet="Datos")["rows"][0]["r"]
        self.assertEqual(len(wide_row), 15)
