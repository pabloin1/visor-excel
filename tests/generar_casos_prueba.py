"""
Script para generar la suite completa de archivos de prueba (Fase 1).
Crea archivos con casos límite de encabezados, codificaciones, tipos mixtos,
ceros a la izquierda, archivos vacíos, corruptos y archivo grande.
"""

import csv
from pathlib import Path
import openpyxl
from openpyxl import Workbook
import pandas as pd
import random

BASE_DIR = Path(__file__).resolve().parent.parent
OUT_DIR = BASE_DIR / "muestras" / "casos_prueba"
OUT_DIR.mkdir(parents=True, exist_ok=True)

print(f"Generando casos de prueba en: {OUT_DIR}")

# -------------------------------------------------------------
# Caso 1: Excel con metadatos antes de encabezados
# -------------------------------------------------------------
wb1 = Workbook()
ws1 = wb1.active
ws1.title = "Reporte"
ws1.append(["GOBIERNO DEL ESTADO DE CHIAPAS - REPORTE ANUAL"])  # Fila 1: Título
ws1.append(["Fecha de emisión: 2026-10-13 | Confidencial"])      # Fila 2: Nota
ws1.append([])                                                    # Fila 3: Vacía
ws1.append(["Codigo", "Nombre", "Departamento", "Presupuesto"])  # Fila 4: Encabezados reales
ws1.append(["EMP-01", "Ana Gomez", "Educacion", 45000])
ws1.append(["EMP-02", "Carlos Ruiz", "Salud", 52000])
ws1.append(["EMP-03", "Beatriz Lara", "Tecnologia", 61000])
ws1.append(["EMP-04", "David Mendez", "Obras", 48000])
wb1.save(OUT_DIR / "01_excel_metadatos_previos.xlsx")
print("- 01_excel_metadatos_previos.xlsx creado")

# -------------------------------------------------------------
# Caso 2: Excel con encabezados duplicados, vacíos y celdas nulas
# -------------------------------------------------------------
wb2 = Workbook()
ws2 = wb2.active
ws2.title = "Inventario"
# Encabezados: ID, Descripcion, vacio, Cantidad, Cantidad (duplicado)
ws2.append(["ID", "Descripcion", "", "Cantidad", "Cantidad"])
ws2.append(["A100", "Monitor 27in", "Lote A", 15, 12])
ws2.append(["A200", "Teclado Mecanico", None, None, 45])
ws2.append(["A300", None, "Lote B", 8, None])
ws2.append(["A400", "Mouse Inalambrico", "Lote C", 30, 28])
wb2.save(OUT_DIR / "02_excel_encabezados_duplicados_y_vacios.xlsx")
print("- 02_excel_encabezados_duplicados_y_vacios.xlsx creado")

# -------------------------------------------------------------
# Caso 3: Excel con identificadores con ceros a la izquierda y tipos mixtos
# -------------------------------------------------------------
wb3 = Workbook()
ws3 = wb3.active
ws3.title = "Padron"
ws3.append(["Clave_Electoral", "Folio_Postal", "Beneficiario", "Importe_Apoyo", "Fecha_Afiliacion"])
# Guardar los ceros como texto explícito en Excel
ws3.cell(row=2, column=1, value="00129482").data_type = 's'
ws3.cell(row=2, column=2, value="02840").data_type = 's'
ws3.cell(row=2, column=3, value="Elena Morales")
ws3.cell(row=2, column=4, value=3500.50)
ws3.cell(row=2, column=5, value="2026-01-15")

ws3.cell(row=3, column=1, value="00482910").data_type = 's'
ws3.cell(row=3, column=2, value="00912").data_type = 's'
ws3.cell(row=3, column=3, value="Jorge Estrada")
ws3.cell(row=3, column=4, value="Pendiente")  # Tipo mixto (texto en columna numerica)
ws3.cell(row=3, column=5, value="2026-02-20")

ws3.cell(row=4, column=1, value="00051283").data_type = 's'
ws3.cell(row=4, column=2, value="07010").data_type = 's'
ws3.cell(row=4, column=3, value="Raquel Dominguez")
ws3.cell(row=4, column=4, value=4200.00)
ws3.cell(row=4, column=5, value="2026-03-05")

ws3.cell(row=5, column=1, value="00993182").data_type = 's'
ws3.cell(row=5, column=2, value="00100").data_type = 's'
ws3.cell(row=5, column=3, value="Martin Orozco")
ws3.cell(row=5, column=4, value="N/A")  # Tipo mixto
ws3.cell(row=5, column=5, value="2026-04-12")

wb3.save(OUT_DIR / "03_excel_tipos_mixtos_y_ceros.xlsx")
print("- 03_excel_tipos_mixtos_y_ceros.xlsx creado")

# -------------------------------------------------------------
# Caso 4: CSV estándar con comas y UTF-8 (con tildes)
# -------------------------------------------------------------
csv_coma = OUT_DIR / "04_csv_coma_utf8.csv"
with open(csv_coma, "w", encoding="utf-8", newline="") as f:
    writer = csv.writer(f)
    writer.writerow(["Municipio", "Region", "Café_Quintales", "Población_Objetivo"])
    writer.writerow(["Tapachula", "Soconusco", 12500, 350000])
    writer.writerow(["San Cristóbal de Las Casas", "Los Altos", 8400, 215000])
    writer.writerow(["Tuxtla Gutiérrez", "Metropolitana", 3200, 600000])
    writer.writerow(["Comitán de Domínguez", "Meseta Comiteca", 6100, 160000])
print("- 04_csv_coma_utf8.csv creado")

# -------------------------------------------------------------
# Caso 5: CSV con punto y coma y codificación Latin-1 (ISO-8859-1)
# -------------------------------------------------------------
csv_latin1 = OUT_DIR / "05_csv_puntoycoma_latin1.csv"
with open(csv_latin1, "w", encoding="latin-1", newline="") as f:
    writer = csv.writer(f, delimiter=";")
    writer.writerow(["Identificador;Folio", "Descripción_Operación", "Monto_Ejecutado"])
    # Nota: la primera columna tiene un encabezado compuesto con punto y coma o delimitado estándar
    writer.writerow(["OP-001", "Adquisición de camiones y equipamiento", 154200.50])
    writer.writerow(["OP-002", "Instalación de paneles solares en región norte", 89400.00])
    writer.writerow(["OP-003", "Capacitación y educación técnica agrícola", 34500.75])
print("- 05_csv_puntoycoma_latin1.csv creado")

# -------------------------------------------------------------
# Caso 6: CSV con tabuladores y UTF-8-SIG (con BOM)
# -------------------------------------------------------------
csv_tab = OUT_DIR / "06_csv_tab_utf8_sig.csv"
with open(csv_tab, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f, delimiter="\t")
    writer.writerow(["Codigo_Postal", "Localidad", "Beneficiarios"])
    writer.writerow(["29000", "Tuxtla Centro", 1450])
    writer.writerow(["29200", "San Cristóbal Barrio", 980])
    writer.writerow(["30700", "Tapachula Puerto", 1200])
print("- 06_csv_tab_utf8_sig.csv creado")

# -------------------------------------------------------------
# Caso 7: Archivos vacíos (0 bytes)
# -------------------------------------------------------------
(OUT_DIR / "07_archivo_vacio.xlsx").write_bytes(b"")
(OUT_DIR / "07_archivo_vacio.csv").write_bytes(b"")
print("- 07_archivo_vacio (.xlsx y .csv) creados")

# -------------------------------------------------------------
# Caso 8: Excel con hojas sin datos / hojas vacías
# -------------------------------------------------------------
wb8 = Workbook()
ws8_1 = wb8.active
ws8_1.title = "HojaConDatos"
ws8_1.append(["Item", "Valor"])
ws8_1.append(["A", 100])
ws8_1.append(["B", 200])

ws8_2 = wb8.create_sheet(title="HojaVaciaTotal")
# Hoja sin filas ni celdas

ws8_3 = wb8.create_sheet(title="HojaSoloEncabezados")
ws8_3.append(["Columna1", "Columna2", "Columna3"])

wb8.save(OUT_DIR / "08_excel_hojas_vacias.xlsx")
print("- 08_excel_hojas_vacias.xlsx creado")

# -------------------------------------------------------------
# Caso 9: Archivo dañado / corrupto (no es un ZIP válido)
# -------------------------------------------------------------
(OUT_DIR / "09_archivo_corrupto.xlsx").write_bytes(b"ESTE NO ES UN EXCEL VALIDO NI UN ARCHIVO ZIP 123456789")
print("- 09_archivo_corrupto.xlsx creado")

# -------------------------------------------------------------
# Caso 10: Archivo grande (35,000 filas) para pruebas de rendimiento
# -------------------------------------------------------------
print("- Generando 10_archivo_grande.csv (35,000 registros)...")
csv_grande = OUT_DIR / "10_archivo_grande.csv"
random.seed(42)
zonas = ["Norte", "Sur", "Centro", "Altos", "Fronteriza", "Selva"]
conceptos = ["Fertilizante", "Semilla Maiz", "Semilla Cafe", "Herramientas", "Riego", "Abono"]

with open(csv_grande, "w", encoding="utf-8", newline="") as f:
    writer = csv.writer(f)
    writer.writerow(["ID_Registro", "Folio_Unico", "Zona", "Concepto", "Cantidad", "Costo_Unitario", "Total", "Fecha"])
    for i in range(1, 35001):
        folio = f"FOL-{i:06d}"
        zona = random.choice(zonas)
        con = random.choice(conceptos)
        cant = random.randint(1, 50)
        costo = round(random.uniform(50.0, 500.0), 2)
        total = round(cant * costo, 2)
        dia = random.randint(1, 28)
        mes = random.randint(1, 12)
        fecha = f"2026-{mes:02d}-{dia:02d}"
        writer.writerow([i, folio, zona, con, cant, costo, total, fecha])

print("- 10_archivo_grande.csv creado exitosamente")
print("\nTodos los archivos de prueba han sido generados en:", OUT_DIR)
