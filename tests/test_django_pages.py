import io
import os
import tempfile
import time
from pathlib import Path

from django.contrib.staticfiles import finders
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook

from documents.models import Document, Version
from documents import pending_uploads
from documents.services import history


def workbook_bytes(amount, extra=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Datos"
    sheet.append(["Clave", "Monto"])
    sheet.append(["001", amount])
    if extra:
        sheet.append(["002", amount + 1])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class DjangoPageTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.settings_override = override_settings(MEDIA_ROOT=root / "uploads", VISOR_CACHE_ROOT=root / "cache")
        self.settings_override.enable()
        self.client = Client(enforce_csrf_checks=True)
        self.client.get("/")

    def tearDown(self):
        self.settings_override.disable()
        self.temp_dir.cleanup()

    def post_file(self, amount, extra=False):
        token = self.client.cookies["csrftoken"].value
        return self.client.post(reverse("documents:upload"), {
            "action": "upload", "csrfmiddlewaretoken": token,
            "file": SimpleUploadedFile("reporte.xlsx", workbook_bytes(amount, extra)),
        })

    def post_named_file(self, filename, amount, *, document_id=None, changed_columns=False):
        token = self.client.cookies["csrftoken"].value
        payload = workbook_bytes(amount)
        if changed_columns:
            book = Workbook()
            sheet = book.active
            sheet.title = "Datos"
            sheet.append(["Clave", "Estado"])
            sheet.append(["001", "Activo"])
            output = io.BytesIO()
            book.save(output)
            payload = output.getvalue()
        return self.client.post(reverse("documents:upload"), {
            "action": "upload", "csrfmiddlewaretoken": token,
            "document_id": document_id or "",
            "file": SimpleUploadedFile(filename, payload),
        })

    def post_action(self, action, **data):
        token = self.client.cookies["csrftoken"].value
        return self.client.post(reverse("documents:upload"), {
            "action": action, "csrfmiddlewaretoken": token, **data,
        })

    def test_server_rendered_library_upload_viewer_and_delete(self):
        library = self.client.get("/")
        self.assertContains(library, "Biblioteca de documentos")
        self.assertNotContains(library, "app.js")
        self.assertIsNone(finders.find("index.html"))
        self.assertIsNotNone(finders.find("documents/app.generated.css"))

        upload_page = self.client.get(reverse("documents:upload"))
        self.assertContains(upload_page, "csrfmiddlewaretoken")
        created = self.post_file(10)
        self.assertEqual(created.status_code, 302)
        version = Version.objects.get()
        self.assertEqual(created["Location"], reverse("documents:version", args=[version.pk]))
        self.assertNotIn(pending_uploads.SESSION_KEY, self.client.session)

        page = self.client.get(created["Location"])
        self.assertContains(page, "reporte.xlsx")
        self.assertContains(page, "Resumen")
        self.assertContains(page, "Distribución de frecuencias: Monto")
        data = self.client.get(created["Location"], {"view": "data", "sheet": "Datos", "q": "001"})
        self.assertContains(data, "1 de 1 filas")
        self.assertContains(data, "001")
        no_rows = self.client.get(created["Location"], {"view": "data", "sheet": "Datos", "filter_0": "999"})
        self.assertContains(no_rows, "Ninguna fila coincide")
        custom_summary = self.client.get(created["Location"], {
            "view": "summary", "sheet": "Datos", "category_col": "0", "measure_col": "1", "stat_op": "avg",
        })
        self.assertContains(custom_summary, "Promedio de Monto por Clave")

        confirmation = self.client.get(reverse("documents:delete", args=[version.pk]))
        self.assertContains(confirmation, "csrfmiddlewaretoken")
        deleted = self.client.post(reverse("documents:delete", args=[version.pk]), {
            "csrfmiddlewaretoken": self.client.cookies["csrftoken"].value,
        })
        self.assertEqual(deleted.status_code, 302)
        self.assertFalse(Version.objects.exists())

    def test_multistep_upload_and_complete_change_history(self):
        self.assertEqual(self.post_file(10).status_code, 302)
        for amount, expected_steps in ((20, 1), (30, 2)):
            conflict = self.post_file(amount)
            self.assertContains(conflict, "Documento existente")
            self.assertContains(conflict, "Guardar como nueva versión")
            setup = self.post_action("new_version")
            self.assertContains(setup, "Identificar registros")
            self.assertContains(setup, 'name="keys_0"')
            invalid = self.post_action("save_keys", keys_0=["Columna inexistente"])
            self.assertContains(invalid, "Seleccione una opción válida")
            self.assertEqual(Version.objects.count(), expected_steps)
            saved = self.post_action("save_keys", keys_0=["Clave"])
            self.assertEqual(saved.status_code, 302, saved.content)
            page = self.client.get(saved["Location"], {"view": "changes", "sheet": "Datos"})
            self.assertEqual(page.status_code, 200)
            self.assertEqual(len(page.context["history"]), expected_steps)
            self.assertContains(page, f"v{expected_steps} → v{expected_steps + 1}")
            self.assertContains(page, "actualizado")
            if expected_steps == 2:
                first_step = next(item for item in page.context["history"] if item["version_number"] == 2)
                earlier = self.client.get(saved["Location"], {
                    "view": "changes", "sheet": "Datos", "change": first_step["id"],
                    "change_sheet": "Datos", "kind": "updated",
                })
                self.assertContains(earlier, "Cambios de v1 a v2")
                self.assertContains(earlier, "Monto")
        self.assertEqual(Version.objects.count(), 3)

    def test_existing_name_can_be_saved_as_new_document(self):
        self.post_file(10)
        self.assertContains(self.post_file(20), "Documento existente")
        saved = self.post_action("new_document")
        self.assertEqual(saved.status_code, 302)
        self.assertEqual(Version.objects.count(), 2)
        self.assertEqual(Version.objects.values("document_id").distinct().count(), 2)

    def test_card_accepts_differently_named_excel_as_new_version(self):
        created = self.post_named_file("padron_inicial.xlsx", 10)
        self.assertEqual(created.status_code, 302)
        document = Document.objects.get()
        library = self.client.get(reverse("documents:library"))
        self.assertContains(library, 'class="library-card-grid"')
        self.assertContains(library, f"?document_id={document.pk}")
        self.assertContains(library, 'library-card-drop-form')
        self.assertContains(library, 'documents/library-drop.js')
        self.assertContains(library, f'data-history-template="library-history-{document.pk}"')
        self.assertContains(library, f'<template id="library-history-{document.pk}">')
        self.assertContains(library, 'id="library-history-dialog"')
        self.assertContains(library, 'documents/library-history.js')

        target_page = self.client.get(reverse("documents:upload"), {"document_id": document.pk})
        self.assertContains(target_page, f"Agregar versión a {document.name}")
        self.assertContains(target_page, f'name="document_id" value="{document.pk}"')
        staged = self.post_named_file("archivo_renombrado.xlsx", 25, document_id=document.pk, changed_columns=True)
        self.assertContains(staged, "Identificar registros")
        self.assertContains(staged, document.name)
        self.assertContains(staged, "Se agregó la columna «Estado»")
        self.assertEqual(Version.objects.count(), 1)
        saved = self.post_action("save_keys", keys_0=["Clave"])
        self.assertEqual(saved.status_code, 302)
        latest = Version.objects.get(version_number=2)
        self.assertEqual(latest.document_id, document.pk)
        self.assertEqual(latest.original_filename, "archivo_renombrado.xlsx")
        self.assertEqual(latest.change_summary["sheets"][0]["updated"], 0)
        self.assertIn("column_added", {change["kind"] for change in latest.change_summary["schema_changes"]})
        self.assertIn("column_removed", {change["kind"] for change in latest.change_summary["schema_changes"]})
        self.assertEqual(Document.objects.count(), 1)
        library = self.client.get(reverse("documents:library"))
        self.assertContains(library, "archivo_renombrado.xlsx")
        self.assertContains(library, "v2")
        by_current_filename = self.client.get(reverse("documents:library"), {"q": "archivo_renombrado"})
        self.assertEqual(by_current_filename.context["total"], 1)
        self.assertContains(by_current_filename, document.name)
        legacy_summary = latest.change_summary.copy()
        legacy_summary.pop("schema_changes")
        Version.objects.filter(pk=latest.pk).update(change_summary=legacy_summary)
        reconstructed = history.get_version(latest.pk)["change_history"][-1]["schema_changes"]
        self.assertIn("column_added", {change["kind"] for change in reconstructed})
        self.assertIn("column_removed", {change["kind"] for change in history.get_version_changes(latest.pk)["schema_changes"]})
        changes = self.client.get(saved["Location"], {"view": "changes", "sheet": "Datos"})
        self.assertContains(changes, "Cambios de hojas y columnas")
        self.assertContains(changes, "Se eliminó la columna «Monto»")

    def test_general_upload_without_name_match_requires_destination(self):
        self.assertEqual(self.post_named_file("padron.xlsx", 10).status_code, 302)
        first = Document.objects.get()
        staged = self.post_named_file("otro_nombre.xlsx", 20)
        self.assertContains(staged, "Elegir destino del archivo")
        self.assertContains(staged, "Añadir a este documento")
        self.assertEqual((Document.objects.count(), Version.objects.count()), (1, 1))
        searched = self.client.get(reverse("documents:upload"), {"stage": "destination", "q": "padron"})
        self.assertContains(searched, first.name)
        selected = self.post_action("select_document", document_id=first.pk)
        self.assertContains(selected, "Identificar registros")
        saved = self.post_action("save_keys", keys_0=["Clave"])
        self.assertEqual(saved.status_code, 302)
        self.assertEqual(Version.objects.get(version_number=2).document_id, first.pk)

        independent = self.post_named_file("separado.xlsx", 30)
        self.assertContains(independent, "Elegir destino del archivo")
        saved_independent = self.post_action("new_document")
        self.assertEqual(saved_independent.status_code, 302)
        self.assertEqual(Document.objects.count(), 2)
        self.assertEqual(Document.objects.get(name="separado").versions.count(), 1)

    def test_invalid_card_target_does_not_create_version(self):
        self.assertEqual(self.client.get(reverse("documents:upload"), {"document_id": "invalido"}).status_code, 404)
        invalid = self.post_named_file("nuevo.xlsx", 10, document_id=9999)
        self.assertContains(invalid, "Seleccione una opción válida")
        self.assertEqual(Version.objects.count(), 0)

    def test_duplicate_file_on_another_card_reports_actual_destination(self):
        first = self.post_named_file("primero.xlsx", 10)
        second = self.post_named_file("segundo.xlsx", 20)
        self.assertContains(second, "Elegir destino del archivo")
        self.assertEqual(self.post_action("new_document").status_code, 302)
        first_version = Version.objects.get(original_filename="primero.xlsx")
        other_document = Document.objects.get(name="segundo")
        duplicate = self.client.post(reverse("documents:upload"), {
            "action": "upload", "csrfmiddlewaretoken": self.client.cookies["csrftoken"].value,
            "document_id": other_document.pk,
            "file": SimpleUploadedFile("renombrado.xlsx", (history.uploads_dir() / first_version.stored_path).read_bytes()),
        })
        self.assertEqual(duplicate.status_code, 302)
        self.assertEqual(duplicate["Location"], reverse("documents:version", args=[first_version.pk]))
        self.assertEqual(other_document.versions.count(), 1)
        page = self.client.get(duplicate["Location"])
        self.assertContains(page, "no se añadió a la tarjeta elegida")

    def test_destination_list_is_paginated_without_losing_pending_file(self):
        for number in range(12):
            history.upload(f"libro_{number:02d}.xlsx", workbook_bytes(number + 1))
        initial = self.post_named_file("sin_coincidencia.xlsx", 100)
        self.assertContains(initial, "Elegir destino del archivo")
        self.assertEqual(len(initial.context["destination_items"]), 10)
        pending_token = self.client.session[pending_uploads.SESSION_KEY]["token"]
        next_page = self.client.get(initial.context["destination_pager"]["next"])
        self.assertEqual(len(next_page.context["destination_items"]), 2)
        self.assertEqual(self.client.session[pending_uploads.SESSION_KEY]["token"], pending_token)
        self.assertEqual(Document.objects.count(), 12)

    def test_card_accepts_version_with_no_shared_columns(self):
        self.assertEqual(self.post_named_file("base.xlsx", 10).status_code, 302)
        document = Document.objects.get()
        book = Workbook()
        sheet = book.active
        sheet.title = "Datos"
        sheet.append(["Referencia", "Estado"])
        sheet.append(["A-001", "Activo"])
        output = io.BytesIO()
        book.save(output)
        token = self.client.cookies["csrftoken"].value
        staged = self.client.post(reverse("documents:upload"), {
            "action": "upload", "csrfmiddlewaretoken": token, "document_id": document.pk,
            "file": SimpleUploadedFile("estructura_nueva.xlsx", output.getvalue()),
        })
        self.assertContains(staged, "Sin columnas compartidas")
        saved = self.post_action("save_keys")
        self.assertEqual(saved.status_code, 302)
        latest = Version.objects.get(document=document, version_number=2)
        change = latest.change_summary["sheets"][0]
        self.assertEqual((change["added"], change["updated"], change["removed"]), (1, 0, 1))
        self.assertEqual(change["status"], "no_common_columns")
        self.assertTrue(latest.change_summary["schema_changes"])

    def test_change_page_opens_first_nonempty_kind(self):
        self.post_file(10)
        self.post_file(10, extra=True)
        self.post_action("new_version")
        saved = self.post_action("save_keys", keys_0=["Clave"])
        self.assertEqual(saved.status_code, 302)
        page = self.client.get(saved["Location"], {"view": "changes", "sheet": "Datos"})
        self.assertEqual(page.context["kind"], "added")
        self.assertContains(page, "002")

    def test_cancel_removes_pending_upload(self):
        self.post_file(10)
        self.assertContains(self.post_file(20), "Documento existente")
        pending = self.client.session[pending_uploads.SESSION_KEY]
        path = pending_uploads.pending_dir() / f'{pending["token"]}.bin'
        self.assertTrue(path.exists())
        response = self.post_action("cancel")
        self.assertEqual(response.status_code, 302)
        self.assertFalse(path.exists())
        self.assertNotIn(pending_uploads.SESSION_KEY, self.client.session)

    def test_reconcile_storage_cleans_old_pending_upload(self):
        self.post_file(10)
        self.post_file(20)
        pending = self.client.session[pending_uploads.SESSION_KEY]
        path = pending_uploads.pending_dir() / f'{pending["token"]}.bin'
        old = time.time() - 48 * 60 * 60
        os.utime(path, (old, old))
        output = io.StringIO()
        call_command("reconcile_storage", stdout=output)
        self.assertIn("Temporal de carga pendiente", output.getvalue())
        self.assertTrue(path.exists())
        call_command("reconcile_storage", "--delete", stdout=io.StringIO())
        self.assertFalse(path.exists())
