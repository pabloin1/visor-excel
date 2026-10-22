"""Genera cuatro snapshots de un padrón para probar el historial completo."""

from argparse import ArgumentParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


OUTPUT_DIR = Path(__file__).resolve().parent / "historial_principal_4_versiones"
FILENAME = "padron_historial_principal.xlsx"
SHEET_NAME = "Padron"
HEADERS = ("Municipio", "Folio", "Programa", "Monto", "Estado")
MUNICIPIOS = (
    "Tuxtla Gutiérrez",
    "Tapachula",
    "Comitán de Domínguez",
    "Palenque",
    "San Cristóbal de Las Casas",
)
PROGRAMAS = ("Alimentación", "Vivienda", "Educación", "Salud")
IDS_BY_VERSION = {
    1: range(1, 101),
    2: range(1, 131),
    3: range(51, 131),
    4: range(51, 161),
}
UPDATED_IDS = {
    2: range(51, 61),
    3: range(61, 69),
    4: range(71, 83),
}
EXPECTED_CHANGES = {
    2: (30, 10, 0),
    3: (0, 8, 50),
    4: (30, 12, 0),
}


def records_for(version: int) -> dict[str, tuple]:
    """Devuelve el contenido completo de la versión, indexado por Folio."""
    if version not in IDS_BY_VERSION:
        raise ValueError(f"Versión desconocida: {version}")
    records = {}
    for number in IDS_BY_VERSION[version]:
        folio = f"AP-{number:04d}"
        amount = 1000 + number * 25
        status = "Activo"
        for edited_in, numbers in UPDATED_IDS.items():
            if edited_in <= version and number in numbers:
                amount += edited_in * 100
                status = f"Revisado en v{edited_in}"
        records[folio] = (
            MUNICIPIOS[(number - 1) % len(MUNICIPIOS)],
            folio,
            PROGRAMAS[(number - 1) % len(PROGRAMAS)],
            amount,
            status,
        )
    return records


def changes_between(previous: dict, current: dict) -> tuple[int, int, int]:
    old_keys, new_keys = set(previous), set(current)
    return (
        len(new_keys - old_keys),
        sum(previous[key] != current[key] for key in old_keys & new_keys),
        len(old_keys - new_keys),
    )


def generate(output_dir: Path = OUTPUT_DIR) -> list[Path]:
    """Escribe los cuatro libros y verifica que contienen los cambios esperados."""
    output_dir = Path(output_dir)
    paths = []
    previous = None
    for version in range(1, 5):
        records = records_for(version)
        if previous is not None:
            actual = changes_between(previous, records)
            if actual != EXPECTED_CHANGES[version]:
                raise AssertionError(f"v{version}: cambios {actual}, esperados {EXPECTED_CHANGES[version]}")

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(HEADERS)
        for row in records.values():
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 25
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="087F72")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for column, width in {"A": 30, "B": 15, "C": 20, "D": 16, "E": 22}.items():
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows(min_row=2, min_col=4, max_col=4):
            row[0].number_format = '"$"#,##0.00'

        path = output_dir / f"v{version}" / FILENAME
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        workbook.close()

        saved = load_workbook(path, read_only=True, data_only=True)
        try:
            saved_sheet = saved[SHEET_NAME]
            actual_rows = list(saved_sheet.values)
            if tuple(actual_rows[0]) != HEADERS or len(actual_rows) - 1 != len(records):
                raise AssertionError(f"v{version}: el libro guardado no coincide con los datos generados")
            if len({row[1] for row in actual_rows[1:]}) != len(records):
                raise AssertionError(f"v{version}: hay folios duplicados")
        finally:
            saved.close()
        paths.append(path)
        previous = records
    return paths


def main() -> None:
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    for version, path in enumerate(generate(args.output_dir), start=1):
        print(f"v{version}: {len(records_for(version))} filas - {path}")
    print("Cambios: v1->v2 +30/10 actualizados; v2->v3 -50/8 actualizados; v3->v4 +30/12 actualizados.")


if __name__ == "__main__":
    main()
